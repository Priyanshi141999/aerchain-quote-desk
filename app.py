"""Quote Desk — Streamlit app. Run: streamlit run app.py"""
from __future__ import annotations

import io
import json
import os
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import streamlit as st

ROOT = Path(__file__).parent
DATA = ROOT / "data"
EXTRACTIONS = DATA / "extractions"

st.set_page_config(page_title="Quote Desk", page_icon="📦", layout="wide")

# keys from Streamlit secrets → environment (the AI layer reads env vars)
for k in ("GEMINI_API_KEY", "ANTHROPIC_API_KEY", "GEMINI_MODEL", "CLAUDE_MODEL"):
    try:
        if k in st.secrets and not os.environ.get(k):
            os.environ[k] = st.secrets[k]
    except Exception:
        pass

import importlib  # noqa: E402
import sys as _sys  # noqa: E402
# Streamlit re-runs app.py on each change but may keep imported helper modules cached. Reload a helper when its
# file changed (and once per server start), looking modules up by NAME so a half-reloaded package can't break the app.
_QD = ["quotedesk.llm", "quotedesk.ingest", "quotedesk.normalize", "quotedesk.checks", "quotedesk.extract",
       "quotedesk.pipeline", "quotedesk.tables", "quotedesk.analyst", "quotedesk.copilot", "quotedesk.updates"]
_seen = getattr(_sys, "_qd_mtimes", None)
_changed = _seen is None
_seen = _seen or {}
for _name in _QD:
    try:
        _m = importlib.import_module(_name)
        _mt = os.path.getmtime(_m.__file__)
        if _changed or _seen.get(_name, _mt) != _mt:
            importlib.reload(_m)
            _changed = True
        _seen[_name] = _mt
    except Exception as _e:  # never let the refresh itself take the app down
        print(f"[app] module refresh skipped for {_name}: {_e}", flush=True)
_sys._qd_mtimes = _seen
from quotedesk import analyst, copilot, llm, updates  # noqa: E402
from quotedesk.normalize import SETTINGS  # noqa: E402
from quotedesk.ingest import read_vendor_folder  # noqa: E402
from quotedesk.pipeline import load_rfq, run_all  # noqa: E402
from quotedesk.tables import attention_queue, line_table, vendor_table  # noqa: E402

st.markdown("""
<style>
.block-container {padding-top: 1.6rem; max-width: 1400px;}
.qd-badge {display:inline-block; padding:2px 9px; border-radius:10px; font-size:0.78rem; font-weight:600;}
.qd-q {background:#E6F4EA; color:#1E6B34;} .qd-c {background:#FFF4D6; color:#7A5600;} .qd-n {background:#FDE7E7; color:#9B1C1C;}
.qd-card {border:1px solid rgba(128,128,128,.25); border-radius:10px; padding:12px 14px; height:100%;}
.qd-muted {color: #6b7280; font-size: 0.85rem;}
.qd-trail {font-family: ui-monospace, monospace; font-size: 0.85rem; background: rgba(128,128,128,.08); padding: 8px 10px; border-radius: 6px;}
</style>
""", unsafe_allow_html=True)

QUAL_BADGE = {"qualified": ("Qualified", "qd-q"), "conditional": ("Conditional", "qd-c"), "not_qualified": ("Not qualified", "qd-n"),
              "disqualified": ("Disqualified by buyer", "qd-n")}
SEV_ICON = {"blocker": "🛑", "warning": "⚠️", "info": "ℹ️"}
VENDOR_FOLDERS = {p.name.split("_")[0]: p for p in sorted((DATA / "vendors").iterdir()) if p.is_dir()}

# ------------------------------------------------------------------ state
ss = st.session_state
ss.setdefault("overrides", {})       # "A:13" -> {value, reason, by, at}
ss.setdefault("audit", [])           # every human decision, in order
ss.setdefault("acks", {})            # flag id -> note
ss.setdefault("chat", [])            # analyst conversation
ss.setdefault("copilot_msgs", [])
ss.setdefault("draft", None)
ss.setdefault("published", False)
ss.setdefault("disputes", {})          # "A:13" -> vendor's explanation of why our reading is wrong
ss.setdefault("buyer_decisions", {})   # vendor key -> buyer's disqualify / accept decisions
ss.setdefault("vendor_updates", {})  # vendor key -> evidence submitted in the portal after the quote

items, questions, rfq_terms, prev = load_rfq(DATA)


def log(action: str, detail: str, by: str = "Priya Raman (buyer)"):
    ss.audit.append({"time": datetime.now().strftime("%d %b %H:%M:%S"), "by": by, "action": action, "detail": detail})


def available_providers():
    out = []
    if os.environ.get("ANTHROPIC_API_KEY"):
        out.append("claude")
    if os.environ.get("GEMINI_API_KEY"):
        out.append("gemini")
    return out


def cached_provider():
    for p in ("claude", "gemini", "mock"):
        if list(EXTRACTIONS.glob(f"raw_*_{p}.json")):
            return p
    return None


def load_results(force_live=False):
    prov = llm.provider()
    has_cache = len(list(EXTRACTIONS.glob(f"raw_*_{prov}.json"))) >= len(VENDOR_FOLDERS)
    if not force_live and not has_cache:
        cp = cached_provider()
        if cp:
            os.environ["LLM_PROVIDER"] = cp
            prov = cp
    res = run_all(DATA, EXTRACTIONS, use_cache=not force_live)
    ss.results, ss.results_provider, ss.results_at = res, prov, datetime.now().strftime("%d %b %H:%M")
    return res


# ------------------------------------------------------------------ sidebar
with st.sidebar:
    st.markdown("### 📦 Quote Desk")
    st.caption("Kaveri Home Appliances · Corrugated packaging RFQ KHA/PKG/RFQ/2026-27/014")
    page = st.radio("Go to", ["1 · Draft RFQ with co-pilot", "2 · Vendor responses", "3 · Comparison",
                              "4 · Needs your attention", "5 · Ask the analyst", "6 · Vendor portal (vendor view)",
                              "7 · Audit log"], label_visibility="collapsed")
    st.divider()
    provs = available_providers()
    if provs:
        choice = st.selectbox("AI model provider", provs, index=0,
                              format_func=lambda p: {"claude": "Claude (Anthropic)", "gemini": "Gemini (Google)"}[p])
        os.environ["LLM_PROVIDER"] = choice
    else:
        st.warning("No API key configured. Showing saved AI readings; live AI features are disabled.")
    st.caption("The AI reads documents and plans analyses. Code does every calculation, so each number is traceable.")

if "results" not in ss:
    with st.spinner("Loading the AI's readings of the vendor responses…"):
        try:
            load_results()
        except Exception as e:
            ss.results = {}
            st.error(f"Could not load vendor readings: {e}")
EFFQ = updates.effective_questions(questions, ss.get("published_questions"))
MUST = [q["id"] for q in EFFQ["questions"] if q["must_pass"]]
R = (updates.apply_buyer_decisions(
        updates.apply(updates.requalify(ss["results"], EFFQ), ss.vendor_updates, EFFQ), ss.buyer_decisions)
     if ss.get("results") else {})
L = line_table(R, items, ss.overrides, ss.disputes) if R else pd.DataFrame()
V = vendor_table(R) if R else pd.DataFrame()


def updates_banner():
    changed = [(k, v["update_summary"]) for k, v in R.items() if v.get("update_summary")]
    if changed:
        txt = " · ".join(f"**{short_name(R[k]['name'])}** {u['before'].replace('_', ' ')} → **{u['after'].replace('_', ' ')}**"
                         if u["before"] != u["after"] else f"**{short_name(R[k]['name'])}** updated ({'; '.join(u['notes'])})"
                         for k, u in changed)
        st.info(f"🔄 Changes since the quotes arrived (vendor updates and your decisions): {txt}. Rankings and flags below reflect them.")


def badge(qual):
    t, c = QUAL_BADGE.get(qual, (qual, "qd-c"))
    return f'<span class="qd-badge {c}">{t}</span>'


def inr(x, dec=2):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "—"
    if dec == 0:
        s = f"{int(round(x)):,}"
        # Indian grouping
        neg = s.startswith("-"); s = s.lstrip("-").replace(",", "")
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:]); head = head[:-2]
        if head:
            parts.insert(0, head)
        return ("-" if neg else "") + "₹" + (",".join(parts) + "," if parts else "") + tail
    return f"₹{x:,.2f}"


def short_name(n: str) -> str:
    """Drop legal suffixes for compact display (full legal names stay everywhere else)."""
    import re as _re
    n = _re.sub(r"\((india)\)|\b(pvt|private|ltd|limited|llp|industries|corrugated boxes)\b\.?", "", n, flags=_re.I)
    return _re.sub(r"\s+", " ", n).strip(" ,.-")


def crore(x):
    return f"₹{x / 1e7:.2f} Cr"


# ================================================================== 1. CO-PILOT
if page.startswith("1"):
    st.header("Draft the RFQ with the co-pilot")
    st.caption("Talk the RFQ into existence. Start from last year's contract, say what changes, and the co-pilot keeps a structured draft and asks what's missing.")
    left, right = st.columns([5, 6], gap="large")
    with left:
        attach = st.toggle("Attach last year's contract (Annapurna FY25-26, 28 lines)", value=True)
        attach_text = (DATA / "buyer/last_year_contract_annapurna_FY25-26.csv").read_text() if attach else None
        up = st.file_uploader("…or attach your own item list (CSV / TXT)", type=["csv", "txt"])
        if up:
            attach_text = up.read().decode("utf-8", "ignore")
        for m in ss.copilot_msgs:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])
        starter = ("Same as last year for the Hosur plant, but add a printed accessory kit mailer (die-cut, E-flute, 4 colours, "
                   "about 60,000 a year) and an E-flute die-cut insert for the induction cooktop (65,000 a year), and upgrade the "
                   "750W mixer box to 22 BF. One-year rate contract from 1 Nov.")
        if not ss.copilot_msgs:
            st.info(f"**Try:** {starter}")
            if st.button("Use this example"):
                ss.pending_copilot = starter
                st.rerun()
        msg = st.chat_input("Tell the co-pilot what you need…") or ss.pop("pending_copilot", None)
        if msg:
            if not available_providers():
                st.error("Add an API key to use the co-pilot live.")
            else:  # show the buyer's message immediately, then work on the reply
                ss.copilot_msgs.append({"role": "user", "content": msg})
                ss.copilot_working = msg
                st.rerun()
        if ss.get("copilot_working"):
            msg = ss.copilot_working
            with st.status("Co-pilot is drafting the RFQ…", expanded=True) as status:
                st.write("Reading last year's contract and applying your changes (usually 20–60 s).")
                try:
                    if len(ss.copilot_msgs) <= 2:
                        ss.copilot_baseline = attach_text  # what code compares every later draft against
                    out = copilot.turn(msg, ss.draft, attach_text if len(ss.copilot_msgs) <= 2 else None, ss.copilot_msgs[:-1],
                                       baseline_text=ss.get("copilot_baseline"), standard_questions=questions["questions"],
                                       on_status=lambda m: status.update(label=f"Co-pilot is drafting… {m}"))
                    ss.draft = out.get("draft") or ss.draft
                    reply = out.get("reply", "")
                    if out.get("open_questions"):
                        reply += "\n\n**Questions for you:**\n" + "\n".join(f"- {q}" for q in out["open_questions"])
                    if out.get("assumptions"):
                        reply += "\n\n**Assumptions I made (change any):**\n" + "\n".join(f"- {a}" for a in out["assumptions"])
                    ss.copilot_msgs.append({"role": "assistant", "content": reply})
                    status.update(label="Draft updated", state="complete")
                except Exception as e:
                    ss.copilot_msgs.append({"role": "assistant", "content": f"⚠️ I couldn't reach the AI just now: {e}"})
                    status.update(label="AI call failed", state="error")
            ss.copilot_working = None
            st.rerun()
    with right:
        d = ss.draft
        if not d:
            st.markdown("#### Draft will appear here")
            st.caption("Line items, questionnaire and terms update after every message.")
        else:
            st.markdown(f"#### {d.get('title', 'RFQ draft')}")
            st.caption(f"{d.get('category','')} · {d.get('delivery_location','')} · {d.get('contract_period','')}")
            dl = pd.DataFrame(d.get("lines", []))
            if not dl.empty:
                n_new = int((dl.get("status") == "new").sum()) if "status" in dl else 0
                n_chg = int((dl.get("status") == "changed").sum()) if "status" in dl else 0
                st.markdown(f"**{len(dl)} line items** · {n_new} new · {n_chg} changed")
                edited = st.data_editor(dl, hide_index=True, width="stretch", height=360, key="draft_editor")
                ss.draft["lines"] = edited.to_dict("records")
            qs_ = d.get("questionnaire", [])
            n_mp = sum(1 for q in qs_ if q.get("must_pass"))
            with st.expander(f"Questionnaire ({len(qs_)} questions · {n_mp} must-pass)", expanded=False):
                st.caption("Tick or untick **must_pass**. A vendor must pass every must-pass question to qualify; "
                           "on publishing, these choices drive qualification everywhere.")
                if qs_:
                    qdf = pd.DataFrame(qs_)[[c for c in ("id", "text", "must_pass", "why") if c in pd.DataFrame(qs_).columns]]
                    qed = st.data_editor(qdf, hide_index=True, width="stretch", key="q_editor",
                                         disabled=[c for c in qdf.columns if c != "must_pass"],
                                         column_config={"must_pass": st.column_config.CheckboxColumn("must_pass"),
                                                        "text": st.column_config.TextColumn(width="large")})
                    ss.draft["questionnaire"] = qed.to_dict("records")
            with st.expander("Commercial terms", expanded=False):
                for t in d.get("terms", []):
                    st.markdown(f"- {t}")
            c1, c2 = st.columns(2)
            if c1.button("📤 Publish RFQ to 5 vendors", type="primary", disabled=ss.published):
                ss.published = True
                ss.published_questions = ss.draft.get("questionnaire") or None
                mp_ = [q["id"] for q in (ss.published_questions or []) if q.get("must_pass")]
                log("RFQ published", f"{len(dl)} lines sent to 5 vendors by email (stubbed channel); must-pass: {', '.join(mp_) or 'standard'}")
                st.rerun()
            if ss.published:
                st.success("Sent by email to Siam Pacific, Deccan Packaging, Vijay Box Works, Sri Murugan and Annapurna. "
                           "Vendors can reply in any format. **Nine days later →** see *2 · Vendor responses*.")
                st.caption("Email is stubbed for the demo. The responses in this demo answer the published packaging RFQ "
                           "KHA/PKG/RFQ/2026-27/014 (30 lines).")
                st.info("Qualification now uses your must-pass questions: **" + ", ".join(MUST) + "**."
                        + (f" New questions not asked in these vendor responses (not scored): "
                           + ", ".join(str(q.get('id')) for q in EFFQ.get('not_scored', [])) if EFFQ.get("not_scored") else ""))

# ================================================================== 2. RESPONSES
elif page.startswith("2"):
    st.header("Vendor responses")
    st.caption("Five vendors, five formats. Nobody was forced into the template. The AI reads each response in whatever shape it arrived.")
    top = st.columns([3, 2])
    with top[0]:
        st.markdown(f"Readings shown: **{ss.get('results_provider', '—')}** · loaded {ss.get('results_at', '—')}")
    with top[1]:
        if st.button("🔄 Re-read all responses with AI now (live)", disabled=not available_providers()):
            with st.spinner("AI is reading 5 vendors' documents (≈1–3 min)…"):
                try:
                    load_results(force_live=True)
                    log("AI re-read", f"All responses re-read live with {llm.provider()}")
                    st.rerun()
                except Exception as e:
                    st.error(f"Live re-read failed: {e}")
    import pdfplumber
    for vk, folder in VENDOR_FOLDERS.items():
        v = R.get(vk)
        name = v["name"] if v else folder.name
        with st.expander(f"**{vk} · {name}**" + (f" — {v['qualification']['overall'].replace('_', ' ')}" if v else ""), expanded=(vk == "D")):
            docs = read_vendor_folder(folder)
            if v:
                priced = sum(1 for l in v["lines"] if l["norm_inr"] is not None)
                bl = sum(1 for f in v["vendor_flags"] if f["severity"] == "blocker")
                st.markdown(f"{badge(v['qualification']['overall'])} &nbsp; **{priced}/30** lines priced · "
                            f"{bl} blocker(s) · read by `{v['meta'].get('model')}` in {v['meta'].get('seconds')} s", unsafe_allow_html=True)
            st.markdown("**Original files as received** (download to check against what the AI read):")
            dl_cols = st.columns(max(len(docs), 1))
            for dc, d in zip(dl_cols, docs):
                dc.download_button(f"⬇️ {d.file}", (folder / d.file).read_bytes(), file_name=d.file,
                                   key=f"dl_{vk}_{d.file}", width="stretch")
            st.markdown("**Preview**")
            tabs = st.tabs([d.file for d in docs])
            for t, d in zip(tabs, docs):
                with t:
                    if d.kind == "image":
                        st.image(d.binary, caption=d.file, width=520)
                    elif d.kind == "pdf":
                        with pdfplumber.open(folder / d.file) as pdf:
                            imgs = [p.to_image(resolution=70).original for p in pdf.pages[:3]]
                        cols = st.columns(len(imgs))
                        for c, im in zip(cols, imgs):
                            c.image(im)
                        if d.meta.get("invisible_text"):
                            st.error("🛑 This PDF contains text invisible to people (white or tiny): "
                                     f"\"{d.meta['invisible_text'][0]['text'][:180]}…\". It is treated as data and ignored.")
                    elif d.kind == "excel":
                        import openpyxl
                        wb_ = openpyxl.load_workbook(folder / d.file, data_only=True)
                        sheet_tabs = st.tabs(wb_.sheetnames)
                        for stab, ws_ in zip(sheet_tabs, wb_.worksheets):
                            with stab:
                                rows_ = [[c for c in r] for r in ws_.iter_rows(values_only=True)]
                                hidden_ = {k for k, dim in ws_.row_dimensions.items() if dim.hidden}
                                df_ = pd.DataFrame([r for i, r in enumerate(rows_, 1) if i not in hidden_]).dropna(how="all").dropna(axis=1, how="all")
                                st.dataframe(df_.astype(str).replace({"None": ""}), width="stretch", height=320, hide_index=True)
                        if d.meta.get("hidden_rows"):
                            st.warning(f"Hidden rows found by software (not shown above, visible only to machines): {d.meta['hidden_rows']}")
                    else:
                        st.text(d.text[:4000])
                        if d.meta.get("hidden_rows"):
                            st.warning(f"Hidden rows found by software: {d.meta['hidden_rows']}")

# ================================================================== 3. COMPARISON
elif page.startswith("3"):
    st.header("Side-by-side comparison")
    if L.empty:
        st.stop()
    updates_banner()
    st.caption(f"Must-pass questions used for qualification: **{', '.join(MUST)}**"
               + (" (as set when the RFQ was published)" if ss.get("published_questions") else " (company standard)"))
    view = st.radio("View", ["💰 Prices", "📋 Questionnaire answers"], horizontal=True, label_visibility="collapsed")
    if view.startswith("📋"):
        ICON = {"yes": "✅", "no": "❌", "partial": "🟡", "pending": "⏳", "not_answered": "➖"}
        st.markdown("Every vendor's answer to every question, as read from their documents (or updated in the portal). "
                    "✅ meets the requirement · ❌ does not · 🟡 partly / needs review · ⏳ promised later · ➖ not answered. "
                    "**Must-pass** questions decide qualification.")
        qtext = {q["id"]: q for q in EFFQ["questions"]}
        rows = []
        for qid, q in qtext.items():
            row = {"Q": qid, "Question": q["text"], "Must-pass": "MUST-PASS" if q["must_pass"] else ""}
            for vk_, v_ in R.items():
                a_ = v_["qualification"]["questions"].get(qid, {})
                ans = (a_.get("answer") or "").strip()
                why = (a_.get("reason") or "").strip()
                stt_ = a_.get("status", "not_answered")
                cell = ans[:80] if ans else why[:80]
                if stt_ != "yes" and ans and why and why not in ans:
                    cell += f" → {why[:70]}"  # claim vs evidence, e.g. "Yes, certified → expired 31 Mar 2026"
                row[short_name(v_["name"])] = f"{ICON.get(stt_, '➖')} {cell}"
            rows.append(row)
        qdf = pd.DataFrame(rows)
        st.dataframe(qdf, hide_index=True, width="stretch", height=430,
                     column_config={"Question": st.column_config.TextColumn(width="medium")})
        # overall line under the table
        st.markdown("**Result:** " + " · ".join(f"{short_name(v_['name'])}: {badge(v_['qualification']['overall'])}"
                                                for v_ in R.values()), unsafe_allow_html=True)
        st.subheader("Look at one question in detail")
        qsel = st.selectbox("Question", list(qtext), format_func=lambda i: f"{i} · {qtext[i]['text'][:90]}")
        st.caption(f"Requirement: {qtext[qsel].get('pass_rule', '')}" + (" · **must-pass**" if qtext[qsel]["must_pass"] else ""))
        for vk_, v_ in R.items():
            a_ = v_["qualification"]["questions"].get(qsel, {})
            raw = next((x for x in (v_["extraction"].get("questionnaire") or []) if x.get("qid") == qsel), {})
            st.markdown(f"{ICON.get(a_.get('status', 'not_answered'), '➖')} **{v_['name']}** — "
                        f"{a_.get('answer') or '_no answer_'}  \n<span class='qd-muted'>Assessment: {a_.get('reason', '')}"
                        f"{' · Source: ' + str(raw.get('source')) if raw.get('source') else ''}</span>", unsafe_allow_html=True)
        st.stop()
    c = st.columns([2, 2, 2, 3])
    basis = c[0].radio("Price basis", ["Basic price", "Landed (incl. freight)"], horizontal=False)
    scope = c[1].radio("Vendors", ["All vendors", "Qualified + conditional", "Qualified only"])
    show_low = c[2].checkbox("Outline low-confidence AI readings", value=True)
    col = "price_inr" if basis == "Basic price" else "landed_inr"

    # vendor cards
    cards = st.columns(len(V))
    for cc, (_, vr) in zip(cards, V.iterrows()):
        q = R[vr["vendor"]]["qualification"]["questions"]
        mp = " ".join(f"{qid}:{'✅' if q[qid]['status'] == 'yes' else '❌' if q[qid]['status'] in ('no', 'not_answered') else '🟡'}"
                      for qid in MUST if qid in q)
        tot = L[(L.vendor == vr["vendor"])]["annual_value_inr"].sum()
        cc.markdown(f"""<div class="qd-card"><b>{vr['vendor']} · {short_name(vr['vendor_name'])}</b><br>{badge(vr['qualification'])}
        <div class="qd-muted" style="margin-top:6px">{mp}<br>{vr['lines_priced']}/30 lines · {crore(tot)} quoted<br>
        {SEV_ICON['blocker']} {vr['blockers']} &nbsp; {SEV_ICON['warning']} {vr['warnings']}</div></div>""", unsafe_allow_html=True)
    st.write("")

    sel = L.copy()
    if scope == "Qualified only":
        sel = sel[sel.qualification == "qualified"]
    elif scope == "Qualified + conditional":
        sel = sel[sel.qualification.isin(["qualified", "conditional"])]
    vendors_in = list(dict.fromkeys(sel.vendor))
    ranked = sel[(~sel.excluded_from_ranking) & sel[col].notna()]
    best = ranked.loc[ranked.groupby("line")[col].idxmin()] if not ranked.empty else ranked
    best_by_line = dict(zip(best.line, best.vendor))

    piv = sel.pivot(index="line", columns="vendor", values=col)
    meta = {(r.vendor, r.line): r for r in sel.itertuples()}
    disp = pd.DataFrame(index=piv.index)
    disp["Item"] = [next(it["name"] for it in items if it["id"] == ln) for ln in piv.index]
    disp["Qty"] = [f"{next(it['annual_qty'] for it in items if it['id'] == ln):,}" for ln in piv.index]
    disp["UoM"] = [next(it["uom"] for it in items if it["id"] == ln) for ln in piv.index]
    for v in vendors_in:
        vals = []
        for ln in piv.index:
            m = meta.get((v, ln))
            x = piv.loc[ln, v] if v in piv else None
            if m is None:
                vals.append("")
            elif x is None or pd.isna(x):
                vals.append("not quoted" if m.status in ("not_quoted", "price_on_request") else
                            ("freight ?" if col == "landed_inr" and m.price_inr is not None else "—"))
            else:
                mark = {"vendor_confirmed": " ✅", "vendor_corrected": " ✅", "buyer_corrected": " ✏️",
                        "buyer_confirmed": " ✏️", "vendor_disputed": " ❓"}.get(getattr(m, "verification", "unverified"), "")
                if m.worst_flag == 0:
                    mark += " 🛑"
                vals.append(f"{x:,.2f}{mark}")
        disp[f"{v}"] = vals
    disp["L1"] = [best_by_line.get(ln, "—") for ln in piv.index]

    def style(df):
        s = pd.DataFrame("", index=df.index, columns=df.columns)
        for ln in df.index:
            for v in vendors_in:
                m = meta.get((v, ln))
                if not m:
                    continue
                if best_by_line.get(ln) == v:
                    s.loc[ln, v] = "background-color: rgba(34,139,34,.18); font-weight:600"
                if m.excluded_from_ranking:
                    s.loc[ln, v] = "color: #b91c1c; text-decoration: line-through"
                elif show_low and m.confidence == "low":
                    s.loc[ln, v] += "; outline: 2px solid #f59e0b"
        return s
    n_ver = int(sel["verification"].astype(str).str.startswith("vendor").sum()) if "verification" in sel else 0
    n_priced = int(sel[col].notna().sum())
    st.markdown(f"**₹ per unit, ex-GST{' + freight' if col == 'landed_inr' else ''}** · green = lowest eligible · "
                "✅ confirmed by vendor · ❓ vendor says our reading is wrong (see their working) · ✏️ set by buyer · "
                "🛑 blocker, needs your decision · no mark = read by AI, "
                "awaiting vendor confirmation · strikethrough = suspected error, excluded from ranking")
    n_disp = int((sel["verification"] == "vendor_disputed").sum()) if "verification" in sel else 0
    st.progress(n_ver / max(n_priced, 1), text=f"{n_ver} of {n_priced} prices confirmed by vendors"
                + (f" · {n_disp} disputed" if n_disp else "")
                + f" · vendors must respond by {SETTINGS['clarification_deadline']:%a %d %b, %H:%M} IST")
    st.dataframe(disp.style.apply(style, axis=None), width="stretch", height=560,
                 column_config={"Item": st.column_config.TextColumn(width="medium")})
    if not best.empty:
        l1_total = (best[col] * best.annual_qty).sum()
        n_l1 = best.line.nunique()
        st.caption(f"Cheapest-per-line across shown vendors: **{crore(l1_total)}** for {n_l1}/30 lines. "
                   f"Lines with no eligible price: {sorted(set(range(1, 31)) - set(best.line)) or 'none'}.")

    st.subheader("Inspect any number")
    ic = st.columns([1, 3])
    vsel = ic[0].selectbox("Vendor", list(R.keys()), format_func=lambda k: f"{k} · {R[k]['name']}")
    lsel = ic[1].selectbox("Line", [it["id"] for it in items], format_func=lambda i: f"{i} · {next(x['name'] for x in items if x['id'] == i)}")
    row = L[(L.vendor == vsel) & (L.line == lsel)].iloc[0]
    a, b = st.columns([3, 2], gap="large")
    with a:
        st.markdown(f"**Vendor wrote:** `{row['as_written'] or '—'}`  \n**Vendor's description:** {row['vendor_description'] or '—'}")
        src = row['source'] or {}
        st.markdown(f"**Found in:** `{src.get('file', '—')}` · {src.get('location', '')}  \n> {src.get('quote', '')}")
        ver_txt = {"vendor_confirmed": "✅ confirmed by the vendor", "vendor_corrected": "✅ corrected by the vendor",
                   "vendor_disputed": "❓ the vendor says this reading is wrong; see their working below",
                   "buyer_corrected": "✏️ set by the buyer", "buyer_confirmed": "✏️ confirmed by the buyer"}.get(row.get('verification', 'unverified'), "read by AI, awaiting vendor confirmation")
        st.markdown(f"**Status:** {ver_txt} · AI reading confidence: **{row['confidence']}**")
        st.markdown("**How we got to the comparable number**")
        st.markdown("<div class='qd-trail'>" + "<br>".join(row['trail'] or ["no price"]) + "</div>", unsafe_allow_html=True)
        if row['freight_note']:
            st.caption(f"Freight: {row['freight_note']}")
        for f in row['flags']:
            st.markdown(f"{SEV_ICON[f['severity']]} **{f['code'].replace('_', ' ')}** — {f['message']}")
        if vsel == "D" and src.get("file", "").endswith(".jpg"):
            st.image(str(VENDOR_FOLDERS["D"] / src["file"]), width=420, caption="Source photo")
    with b:
        st.markdown("**Correct or confirm this value**")
        with st.form(f"fix_{vsel}_{lsel}"):
            nv = st.number_input("₹ per unit, ex-GST", value=float(row['price_inr']) if row['price_inr'] is not None and not pd.isna(row['price_inr']) else 0.0, step=0.01, format="%.2f")
            reason = st.text_input("Reason (required)", placeholder="e.g. Confirmed with vendor by phone, 6 Oct")
            if st.form_submit_button("Save correction"):
                if not reason.strip():
                    st.error("A reason is required. It goes into the audit log.")
                else:
                    ss.overrides[f"{vsel}:{lsel}"] = {"value": nv, "reason": reason, "by": "Priya Raman (buyer)", "source": "buyer",
                                                     "at": datetime.now().isoformat()}
                    log("Value corrected", f"{vsel} line {lsel}: {row['price_inr']} → {nv:.2f} ({reason})")
                    st.rerun()

# ================================================================== 4. ATTENTION
elif page.startswith("4"):
    st.header("Needs your attention")
    updates_banner()
    st.caption("Everything the system was unsure about, or that could change the award, worst first. Nothing here was silently decided for you.")
    Q = attention_queue(R, L)
    done = sum(1 for i, x in enumerate(Q) if f"{x['vendor']}|{x['line']}|{x['code']}|{i}" in ss.acks)
    st.progress(done / max(len(Q), 1), text=f"{done} of {len(Q)} reviewed")
    sev_f = st.multiselect("Show", ["blocker", "warning"], default=["blocker", "warning"],
                           format_func=lambda s: f"{SEV_ICON[s]} {s}s")
    # ---- buyer decisions: judgement calls only the buyer can make
    decidable = [x for x in Q if x["line"] is None and x["code"] in updates.BUYER_DECIDES
                 and R[x["vendor"]]["qualification"]["overall"] != "disqualified"]
    if decidable:
        st.subheader("Your decisions")
        st.caption("These can't be fixed by the vendor. Decide, with a reason; it goes into the audit log.")
        for x in decidable:
            vk_ = x["vendor"]
            with st.container(border=True):
                st.markdown(f"🛑 **{x['vendor_name']}** — **{updates.BUYER_DECIDES[x['code']]}**: {x['message']}")
                reason = st.text_input("Reason (required)", key=f"dec_r_{vk_}_{x['code']}",
                                       placeholder="e.g. Deadline was clearly communicated; late bids are not accepted")
                b1, b2, _ = st.columns([2, 2, 3])
                if b1.button("⛔ Disqualify vendor", key=f"dq_{vk_}_{x['code']}"):
                    if not reason.strip():
                        st.error("Please give a reason.")
                    else:
                        ss.buyer_decisions.setdefault(vk_, {})["disqualify"] = {"reason": reason, "by": "Priya Raman (buyer)",
                                                                                "at": datetime.now().isoformat(), "flag": x["code"]}
                        log("Vendor disqualified", f"{vk_} ({updates.BUYER_DECIDES[x['code']]}): {reason}")
                        st.rerun()
                if b2.button("✅ Accept and keep in evaluation", key=f"ac_{vk_}_{x['code']}"):
                    if not reason.strip():
                        st.error("Please give a reason.")
                    else:
                        ss.buyer_decisions.setdefault(vk_, {}).setdefault("accept", {})[x["code"]] = {
                            "reason": reason, "by": "Priya Raman (buyer)", "at": datetime.now().isoformat()}
                        log("Buyer accepted", f"{vk_} ({updates.BUYER_DECIDES[x['code']]}): {reason}")
                        st.rerun()
    dq_list = [(k, d["disqualify"]) for k, d in ss.buyer_decisions.items() if d.get("disqualify")]
    if dq_list:
        st.subheader("Disqualified vendors")
        for k, d in dq_list:
            c1_, c2_ = st.columns([5, 1])
            c1_.markdown(f"⛔ **{R[k]['name']}**: {d['reason']} <span class='qd-muted'>({d['at'][:16].replace('T', ' ')})</span>", unsafe_allow_html=True)
            if c2_.button("Reinstate", key=f"rein_{k}"):
                ss.buyer_decisions[k].pop("disqualify", None)
                log("Vendor reinstated", f"{k}")
                st.rerun()
    st.subheader("Everything else")
    for qi, x in enumerate(Q):
        if x["severity"] not in sev_f:
            continue
        key = f"{x['vendor']}|{x['line']}|{x['code']}|{qi}"
        acked = key in ss.acks
        c1, c2 = st.columns([8, 2])
        where = f" · line {x['line']}" if x["line"] else ""
        c1.markdown(f"{SEV_ICON[x['severity']]} **{x['vendor']} · {x['vendor_name']}**{where} — "
                    f"**{x['code'].replace('_', ' ')}**: {x['message']}" + ("  ✅ *reviewed*" if acked else ""))
        if not acked and c2.button("Mark reviewed", key="ack" + key):
            ss.acks[key] = datetime.now().isoformat()
            log("Flag reviewed", f"{x['vendor']} line {x['line'] or '-'}: {x['code']}")
            st.rerun()

# ================================================================== 5. ANALYST
elif page.startswith("5"):
    st.header("Ask the analyst")
    st.caption("Plain-English questions over the whole comparison. The AI writes the calculation, code runs it on the extracted data, and the AI explains only what was computed.")
    examples = {
        "Split it: cheapest per line, qualified vendors only": "What if we split it, cheapest per line, but only among vendors who cleared the quality questionnaire?",
        "Who's cheapest overall, like for like?": "Who is cheapest overall on a like-for-like basis, and where is freight unknown?",
        "Everything to Deccan for 5% off vs the best split": "If we gave everything to Deccan to get their 5% discount, how would that compare with the best split?",
        "Which lines have one or no eligible quote?": "Which lines have only one or no eligible quote among qualified and conditional vendors?",
        "Chart the price spread for 5-ply boxes": "Chart the price spread per line across vendors for the 5-ply boxes.",
        "How do this year's prices compare with last year?": "How do this year's prices compare with last year's contract for the same items?",
        "Export a line-wise award recommendation": "Export a line-wise award recommendation to Excel.",
    }
    st.caption("Try one of these, or type your own below:")
    ec = st.columns(2)
    for i, (label, full) in enumerate(examples.items()):
        if ec[i % 2].button(label, key=f"ex{i}", width="stretch", help=full):
            ss.pending_q = full
            st.rerun()
    for i, h in enumerate(ss.chat):
        with st.chat_message("user"):
            st.markdown(h["q"])
        with st.chat_message("assistant"):
            st.markdown(h["answer"])
            nc = h.get("number_check") or {}
            if nc.get("checked"):
                if nc["unverified"]:
                    st.warning(f"⚠️ {len(nc['unverified'])} of {nc['checked']} figures could not be traced to the calculation: "
                               + ", ".join(nc["unverified"]) + ". Treat them with caution.")
                else:
                    st.caption(f"✓ All {nc['checked']} figures in this answer trace back to the computed results.")
            if h.get("caveats"):
                st.markdown("**Caveats:** " + " · ".join(h["caveats"]))
            if h.get("fig") is not None:
                st.plotly_chart(h["fig"], width="stretch", key=f"fig{i}")
            if isinstance(h.get("table"), pd.DataFrame):
                st.dataframe(h["table"], width="stretch", hide_index=True)
            exp = h.get("export") if isinstance(h.get("export"), pd.DataFrame) else (h.get("table") if isinstance(h.get("table"), pd.DataFrame) else None)
            if exp is not None:
                buf = io.BytesIO(); exp.to_excel(buf, index=False)
                st.download_button("⬇️ Download as Excel", buf.getvalue(), file_name=f"analysis_{i + 1}.xlsx", key=f"dl{i}")
            if h.get("plan"):
                with st.expander("How I worked this out"):
                    st.markdown(f"**Interpretation:** {h['plan'].get('interpretation')}")
                    for a_ in h["plan"].get("assumptions", []):
                        st.markdown(f"- {a_}")
                    st.code(h["plan"].get("code", ""), language="python")
                    if h.get("error"):
                        st.error(h["error"])
            if h.get("followups"):
                st.caption("You could ask next: " + " · ".join(h["followups"]))
    q = st.chat_input("Ask anything about the quotes…") or ss.pop("pending_q", None)
    if q:
        if not available_providers():
            st.error("Add an API key to ask questions live.")
        else:
            ss.analyst_working = q
            st.rerun()
    if ss.get("analyst_working"):
        q = ss.analyst_working
        with st.chat_message("user"):
            st.markdown(q)
        with st.status("Working on it…", expanded=True) as status:
            st.write("1. AI plans the calculation → 2. code runs it on the extracted quotes → 3. AI writes up only what was computed.")
            try:
                facts = []
                for vk_, v_ in R.items():
                    qq = v_["qualification"]["questions"]
                    must = "; ".join(f"{qid} {qq[qid]['status']}: {qq[qid]['reason'][:110]}" for qid in MUST if qid in qq)
                    blk = "; ".join(f["message"][:120] for f in v_["vendor_flags"] if f["severity"] == "blocker")
                    facts.append(f"{vk_} {v_['name']}: {v_['qualification']['overall']} | must-pass: {must} | blockers: {blk or 'none'}")
                out = analyst.ask(q, L, V, items, ss.chat, on_status=lambda m: status.update(label=f"Working on it… {m}"),
                                  vendor_facts="\n".join(facts),
                                  last_year=pd.read_csv(DATA / "buyer/last_year_contract_annapurna_FY25-26.csv"))
                res = out["result"]
                if isinstance(res, dict):
                    frames = [v for v in res.values() if isinstance(v, (pd.DataFrame, pd.Series))]
                    res = max(frames, key=len) if frames else None
                tbl = res if isinstance(res, pd.DataFrame) else (res.reset_index() if isinstance(res, pd.Series) else None)
                ss.chat.append({"q": q, "answer": out["answer"].get("answer_markdown", ""), "caveats": out["answer"].get("caveats", []),
                                "followups": out["answer"].get("followups", []), "plan": out["plan"], "fig": out["fig"],
                                "table": tbl, "export": out.get("export"), "error": out.get("error"),
                                "number_check": out.get("number_check")})
                log("Analyst question", q)
                status.update(label="Done", state="complete")
            except Exception as e:
                ss.chat.append({"q": q, "answer": f"⚠️ I couldn't reach the AI just now: {e}", "caveats": [], "followups": [],
                                "plan": {}, "fig": None, "table": None, "export": None, "error": None, "number_check": None})
                status.update(label="AI call failed", state="error")
        ss.analyst_working = None
        st.rerun()

# ================================================================== 6. VENDOR PORTAL
elif page.startswith("6"):
    st.header("Vendor portal: confirm what we read")
    st.caption("Each vendor gets a private link showing how their quote was read. They confirm or correct their own numbers, "
               "so the person who knows what they meant does the checking. A number is only trusted once its vendor confirms it.")
    vk = st.selectbox("Viewing as", list(R.keys()), format_func=lambda k: f"{R[k]['name']} (vendor)")
    v = R[vk]
    st.markdown(f"Dear **{v['name']}**, thank you for your quote against **RFQ KHA/PKG/RFQ/2026-27/014**.")
    tab_elig, tab_lines = st.tabs(["① Your eligibility", "② Confirm your prices"])

    # ---------------- ① eligibility: ask for exactly what blocks this vendor, check the evidence, re-qualify
    with tab_elig:
        qual = v["qualification"]["overall"]
        st.markdown(f"**Current status:** {badge(qual)}", unsafe_allow_html=True)
        if v.get("update_summary"):
            u = v["update_summary"]
            st.success(f"Updates received {u['at']}: {'; '.join(u['notes'])}. Status: {u['before'].replace('_', ' ')} → {u['after'].replace('_', ' ')}.")
        reqs = updates.requirements(v, EFFQ)
        fixable = [r for r in reqs if r["kind"] != "buyer_decision"]
        if qual == "disqualified":
            st.error(f"This quote is no longer under consideration. {v['qualification'].get('reason', '')}")
        elif not fixable:
            st.success("Nothing further is needed from you for eligibility.")
        else:
            must_fix = [r for r in fixable if r["id"] in MUST]
            if must_fix:
                st.markdown(f"To be **eligible for award**, please resolve the following **{len(fixable)} item(s)** "
                            f"({len(must_fix)} must-pass):")
            else:
                st.markdown(f"Your quote meets the must-pass requirements. To **finalise** it, please resolve "
                            f"**{len(fixable)} commercial item(s)**:")
            with st.form(f"elig_{vk}"):
                answers = {}
                for r in fixable:
                    st.markdown(f"**{r['title']}**  \n<span class='qd-muted'>Why: {r['why']}</span>", unsafe_allow_html=True)
                    if r["kind"] == "certificate":
                        answers[r["id"]] = st.file_uploader(r["ask"], type=["pdf", "png", "jpg", "jpeg"], key=f"cert_{vk}")
                    elif r["kind"] == "commitment":
                        answers[r["id"]] = st.checkbox(r["ask"], key=f"chk_{vk}_{r['id']}")
                    elif r["kind"] == "validity":
                        answers[r["id"]] = st.date_input(r["ask"], value=SETTINGS["deadline"].date() + timedelta(days=90),
                                                         min_value=SETTINGS["deadline"].date(), key=f"val_{vk}")
                    elif r["kind"] == "freight":
                        answers[r["id"]] = st.radio(r["ask"], ["Prices include freight to Hosur (FOR Hosur)", "Freight is extra"],
                                                    index=None, key=f"fr_{vk}")
                    st.divider()
                _dl = SETTINGS["clarification_deadline"]
                _open = datetime.now(_dl.tzinfo) <= _dl
                st.caption(f"Respond by {_dl:%d %b %Y, %H:%M} IST — the same deadline applies to every vendor.")
                submitted = st.form_submit_button("Submit to buyer", type="primary", disabled=not _open)
            if submitted:
                up = dict(ss.vendor_updates.get(vk, {}))
                up["at"] = datetime.now().strftime("%Y-%m-%d")
                msgs = []
                for r in fixable:
                    a = answers.get(r["id"])
                    if r["kind"] == "certificate" and a is not None:
                        mime = {"pdf": "application/pdf", "png": "image/png"}.get(a.name.rsplit(".", 1)[-1].lower(), "image/jpeg")
                        with st.status("Reading your certificate…", expanded=False) as stt:
                            try:
                                cert = updates.read_certificate(a.getvalue(), mime, on_status=lambda m: stt.update(label=f"Reading your certificate… {m}"))
                                ok, why = updates.validate_certificate(cert, v["name"])
                            except Exception as e:
                                ok, why, cert = False, f"Could not read the file: {e}", {}
                        if ok:
                            up["certificate"] = {"accepted": True, "data": cert, "file": a.name}
                            msgs.append(("ok", f"Certificate accepted: {why}"))
                        else:
                            msgs.append(("err", f"Certificate not accepted: {why}"))
                        log("Vendor certificate " + ("accepted" if ok else "rejected"), f"{vk} {a.name}: {why}", by=f"{v['name']} (vendor)")
                    elif r["kind"] == "commitment" and a:
                        up[r["id"]] = True
                        msgs.append(("ok", f"{r['title']}: confirmed"))
                        log("Vendor commitment", f"{vk} {r['id']}: {r['ask']}", by=f"{v['name']} (vendor)")
                    elif r["kind"] == "validity" and a:
                        if a >= SETTINGS["deadline"].date() + timedelta(days=90):
                            up["validity_until"] = a.isoformat()
                            msgs.append(("ok", f"Validity extended to {a:%d %b %Y}"))
                            log("Vendor validity extension", f"{vk}: valid until {a:%d %b %Y}", by=f"{v['name']} (vendor)")
                        else:
                            msgs.append(("err", "Validity must run at least 90 days from the RFQ deadline."))
                    elif r["kind"] == "freight" and a:
                        up["freight"] = "included" if a.startswith("Prices include") else "extra"
                        msgs.append(("ok", f"Freight clarified: {up['freight']}"))
                        log("Vendor clarification", f"{vk}: freight {up['freight']}", by=f"{v['name']} (vendor)")
                ss.vendor_updates[vk] = up
                ss.portal_msgs = msgs
                st.rerun()
            for kind, m in ss.pop("portal_msgs", []):
                (st.success if kind == "ok" else st.error)(m)
        others = [r for r in reqs if r["kind"] == "buyer_decision"]
        if others:
            st.caption("Also under the buyer's review (no action needed from you): " + " ".join(r["why"] for r in others))

    with tab_lines:
        st.markdown("Below is how we read every line. For each one, tell us whether **our reading is correct**. "
                    "If it isn't, **show how your quoted figure converts to a price per our unit of measure** "
                    "(e.g. \"₹38/kg × 0.21 kg actual box weight = ₹7.98 per box\"). "
                    "This step verifies your existing quote; **prices cannot be changed here**. "
                    "Lines that most need your check are at the top.")
        DISPUTE_COL = "If not, show how your quote converts to ₹ per unit"
        rows = []
        for l in v["lines"]:
            it = next(x for x in items if x["id"] == l["rfq_line"])
            key = f"{vk}:{it['id']}"
            ov, dp = ss.overrides.get(key), ss.disputes.get(key)
            doubts = [f["message"] for f in l["flags"] if f["severity"] in ("blocker", "warning")]
            if l["confidence"] == "low":
                doubts.insert(0, "We were not sure we read this correctly.")
            sev = 0 if any(f["severity"] == "blocker" for f in l["flags"]) else (1 if l["confidence"] == "low" else (2 if doubts else 3))
            steps = [t for t in (l.get("trail") or [])[1:-1]]
            read = l["norm_inr"]
            done = ov is not None and ov.get("source") == "vendor"
            rows.append({
                "Line": it["id"], "Item": it["name"], "Qty": it["annual_qty"], "UoM": it["uom"],
                "You wrote": l.get("as_written") or ("not quoted" if read is None else "—"),
                "How we converted it": " → ".join(steps) if steps else "no conversion needed",
                "We read it as (₹ per unit, ex-GST)": read,
                "Is our reading correct?": "Yes" if done else ("No" if dp else None),
                DISPUTE_COL: dp["explanation"] if dp else "",
                "Why please check": (doubts[0][:140] if doubts else ""),
                "Status": "✅ confirmed" if done else ("❓ sent to buyer for review" if dp else "awaiting your answer"),
                "_priority": sev, "_value": -((read or 0) * it["annual_qty"]),
            })
        df_p = pd.DataFrame(rows).sort_values(["_priority", "_value", "Line"]).drop(columns=["_priority", "_value"])
        dl = SETTINGS["clarification_deadline"]
        now = datetime.now(dl.tzinfo)
        open_ = now <= dl
        left = dl - now
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Lines in your quote", len(df_p))
        c2.metric("Need a closer look", int((df_p["Why please check"] != "").sum()))
        c3.metric("Confirmed so far", int((df_p["Status"] == "✅ confirmed").sum()))
        c4.metric("Respond by", f"{dl:%d %b, %H:%M}",
                  delta=(f"{left.days} d {left.seconds // 3600} h left" if open_ else "closed"),
                  delta_color="off" if open_ else "inverse")
        if not open_:
            st.error(f"The clarification window closed on {dl:%d %b %Y, %H:%M} IST. Responses are no longer accepted.")
        edited = st.data_editor(
            df_p, hide_index=True, width="stretch", height=460, key=f"portal_{vk}",
            disabled=[c for c in df_p.columns if c not in ("Is our reading correct?", DISPUTE_COL)] if open_ else True,
            column_config={
                "Is our reading correct?": st.column_config.SelectboxColumn(options=["Yes", "No"], required=False,
                    help="Yes confirms our reading. No sends your working to the buyer for review."),
                DISPUTE_COL: st.column_config.TextColumn(width="large", max_chars=300,
                    help="Show the steps from your quoted figure to a price per our unit, excluding GST."),
                "We read it as (₹ per unit, ex-GST)": st.column_config.NumberColumn(format="%.2f"),
                "How we converted it": st.column_config.TextColumn(width="large"),
                "Why please check": st.column_config.TextColumn(width="large"),
            })
        if st.button("📨 Send my answers to the buyer", type="primary", disabled=not open_):
            n_conf = n_disp = 0
            missing = []
            for _, r in edited.iterrows():
                ans, key = r["Is our reading correct?"], f"{vk}:{int(r['Line'])}"
                read = r["We read it as (₹ per unit, ex-GST)"]
                if ans == "Yes" and read is not None and not pd.isna(read):
                    ss.overrides[key] = {"value": float(read), "source": "vendor", "by": f"{v['name']} (vendor)",
                                         "reason": "confirmed by vendor via portal", "at": datetime.now().isoformat()}
                    ss.disputes.pop(key, None)
                    n_conf += 1
                elif ans == "No":
                    why = str(r[DISPUTE_COL] or "").strip()
                    if not why:
                        missing.append(int(r["Line"]))
                        continue
                    if ss.disputes.get(key, {}).get("explanation") != why:
                        ss.disputes[key] = {"explanation": why, "by": f"{v['name']} (vendor)", "at": datetime.now().isoformat()}
                        log("Vendor disputed a reading", f"{vk} line {int(r['Line'])}: \"{why}\"", by=f"{v['name']} (vendor)")
                    n_disp += 1
            if n_conf or n_disp:
                log("Vendor confirmation", f"{vk}: {n_conf} lines confirmed, {n_disp} disputed with working",
                    by=f"{v['name']} (vendor)")
            ss.portal_line_msg = (n_conf, n_disp, missing)
            st.rerun()
        if ss.get("portal_line_msg"):
            n_conf, n_disp, missing = ss.pop("portal_line_msg")
            st.success(f"Sent: {n_conf} readings confirmed, {n_disp} sent to the buyer with your working.")
            if missing:
                st.error(f"Lines {missing}: you answered No but didn't show your working, so they weren't sent.")

# ================================================================== 7. AUDIT LOG
elif page.startswith("7"):
    st.header("Audit log")
    st.caption("Every human decision is recorded: corrections, vendor confirmations, reviewed flags and questions asked. That's what makes an award defensible.")
    st.dataframe(pd.DataFrame(ss.audit or [{"time": "—", "by": "—", "action": "No actions yet", "detail": ""}]),
                 hide_index=True, width="stretch")
