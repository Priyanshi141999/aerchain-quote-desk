import csv, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from spec import ITEMS, BUYER, last_year_price

OUT = "/home/claude/aerchain/dataset/buyer"
os.makedirs(OUT, exist_ok=True)

# 1. Line items --------------------------------------------------------------
cols = ["line_no", "item", "form", "board_grade", "construction", "dimensions", "bursting_factor_bf",
        "print_colours", "annual_qty", "uom", "approx_weight_kg_per_uom"]
with open(f"{OUT}/rfq_line_items.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(cols)
    for it in ITEMS:
        w.writerow([it["id"], it["name"], it["form"], it["board"], it["construction"], it["dims"],
                    it["bf"] or "", it["print_colours"], it["annual_qty"], it["uom"], it["weight_kg"]])

# 2. Questionnaire -------------------------------------------------------------
QUESTIONS = [
    {"id": "Q1", "text": "Is your manufacturing unit ISO 9001:2015 certified? Attach a valid certificate.", "type": "yes_no+document", "must_pass": True,
     "pass_rule": "Yes, with a certificate valid on the RFQ deadline and issued to the bidding legal entity"},
    {"id": "Q2", "text": "Will you provide lot-wise test reports for Box Compression Test (BCT), bursting strength and moisture content?", "type": "yes_no", "must_pass": True,
     "pass_rule": "Yes (unconditional)"},
    {"id": "Q3", "text": "Monthly corrugation capacity (metric tonnes) and current utilisation (%).", "type": "number", "must_pass": False,
     "pass_rule": "Info only; flag if our monthly volume exceeds 25% of free capacity"},
    {"id": "Q4", "text": "Standard lead time from PO to delivery at Hosur (days).", "type": "number", "must_pass": False,
     "pass_rule": "Preferred <= 21 days"},
    {"id": "Q5", "text": "Kraft paper source: virgin / recycled / mixed. Is FSC certified paper available on request?", "type": "text", "must_pass": False,
     "pass_rule": "Info only"},
    {"id": "Q6", "text": "Do you have in-house flexo printing up to 4 colours?", "type": "yes_no", "must_pass": False,
     "pass_rule": "Required for lines 3 and 20 (4-colour print)"},
    {"id": "Q7", "text": "Years of experience supplying consumer durables / appliance OEMs. Name two current customers.", "type": "text", "must_pass": False,
     "pass_rule": "Info only"},
    {"id": "Q8", "text": "Do you accept our standard payment terms of 60 days from invoice?", "type": "yes_no", "must_pass": True,
     "pass_rule": "Yes"},
    {"id": "Q9", "text": "GSTIN of the billing entity.", "type": "text", "must_pass": False,
     "pass_rule": "Must be present for PO creation"},
    {"id": "Q10", "text": "Will you replace rejected lots within 7 days at your cost, including freight?", "type": "yes_no", "must_pass": False,
     "pass_rule": "Preferred yes"},
]
with open(f"{OUT}/questionnaire.json", "w") as f:
    json.dump({"rfq_no": BUYER["rfq_no"], "questions": QUESTIONS,
               "qualification_rule": "A vendor is QUALIFIED only if every must-pass question (Q1, Q2, Q8) passes with evidence."}, f, indent=2)

# 3. Last year's contract (incumbent) ------------------------------------------------
with open(f"{OUT}/last_year_contract_annapurna_FY25-26.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["contract_no", "vendor", "line_no", "item", "form", "board_grade", "construction", "dimensions",
                "bursting_factor_bf", "print_colours", "annual_qty", "uom", "price_inr_ex_gst", "freight_terms", "valid_till"])
    for it in ITEMS:
        p = last_year_price(it)
        if p is None:
            continue
        bf = 20 if it["id"] == 5 else it["bf"]
        w.writerow(["KHA/PKG/CT/2025-26/003", "Annapurna Packers", it["id"], it["name"], it["form"], it["board"], it["construction"],
                    it["dims"], bf or "", it["print_colours"], it["annual_qty"], it["uom"], p, "FOR Hosur (included)", "2026-10-31"])

# 4. Quote template sent to vendors --------------------------------------------------
wb = Workbook(); ws = wb.active; ws.title = "Price Bid"
ws["A1"] = f"{BUYER['name']} | RFQ {BUYER['rfq_no']} | Price bid template"; ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = "Please quote basic price per UoM in INR, excluding GST. State freight separately. Do not change line numbers."
hdr = ["Line", "Item", "Board", "Dimensions", "BF", "Colours", "Annual qty", "UoM", "Basic price INR / UoM", "Freight INR / UoM", "Lead time (days)", "Remarks"]
ws.append([]); ws.append(hdr)
for c in ws[4]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(wrap_text=True)
for it in ITEMS:
    ws.append([it["id"], it["name"], it["board"], it["dims"], it["bf"] or "", it["print_colours"], it["annual_qty"], it["uom"], None, None, None, None])
for col, wdt in zip("ABCDEFGHIJKL", [6, 40, 11, 44, 5, 8, 11, 7, 14, 12, 10, 24]):
    ws.column_dimensions[col].width = wdt
wb.save(f"{OUT}/vendor_quote_template.xlsx")

# 5. RFQ document (terms) -------------------------------------------------------------
terms = f"""# Request for Quotation - Corrugated Packaging (Annual Rate Contract)

**RFQ No:** {BUYER['rfq_no']}
**Buyer:** {BUYER['name']} - {BUYER['plant']}
**Issued:** {BUYER['issued']}  |  **Response deadline:** {BUYER['deadline']}
**Contact:** {BUYER['buyer_name']}, {BUYER['buyer_title']} ({BUYER['buyer_email']})
**Contract period:** {BUYER['contract_period']}

## Scope
Supply of 30 corrugated packaging items (boxes, partitions, inserts, pads, rolls) for our Hosur plant,
as an annual rate contract with monthly call-off schedules. Indicative annual quantities are listed in
the line-item sheet; actual call-offs may vary +/-20%.

## Commercial terms
1. Quote **basic price per UoM in INR, excluding GST**. Show GST rate separately.
2. Prices to be **FOR Hosur plant**. If freight is extra, state the basis (per trip / per box / at actuals).
3. Prices to be **firm for the full contract period**. State any escalation clause explicitly.
4. Payment terms: **60 days from invoice** (see Q8).
5. One-time costs (printing plates, cutting dies) must be quoted separately.
6. Quote validity: minimum **90 days** from the response deadline.
7. Partial quotes are accepted; clearly mark lines not quoted.
8. Any deviation from specification (board grade, BF, ply) must be declared line-wise.

## Quality
- Board grades and BF as specified. Lot-wise BCT, burst and moisture reports (see Q2).
- Rejected lots to be replaced within 7 days.

## Evaluation
Techno-commercial. Only vendors passing must-pass questions (Q1, Q2, Q8) are eligible for award.
Commercial comparison on landed cost (basic + freight) per UoM x annual quantity. Split award by line is possible.
"""
with open(f"{OUT}/rfq_document.md", "w") as f:
    f.write(terms)

json.dump(ITEMS, open(f"{OUT}/rfq_line_items.json", "w"), indent=2)
print("buyer files written:", os.listdir(OUT))
