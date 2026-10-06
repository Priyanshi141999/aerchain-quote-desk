import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from spec import ITEMS, BUYER
from truth import TRUTH

OUT = "/home/claude/aerchain/dataset/vendors/C_vijay"
os.makedirs(OUT, exist_ok=True)
T = TRUTH["C"]; I = {it["id"]: it for it in ITEMS}

def r(i):
    return f"Rs. {T[i]['raw']:.2f}"

def indian(n):
    s = str(n)
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:]); head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts) + "," + tail

d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = h.add_run("VIJAY BOX WORKS"); run.bold = True; run.font.size = Pt(18); run.font.color.rgb = RGBColor(0x1A, 0x5E, 0x20)
p = d.add_paragraph("No. 7/2, Bagalur Road, Mookandapalli, Hosur - 635126  |  Ph: 94433 50219  |  vijayboxworks@gmail.com")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size = Pt(9)
p = d.add_paragraph("Udyam Reg. No. UDYAM-TN-12-0048871 (Micro Enterprise)  |  GSTIN 33ABQPV7731D1Z9"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size = Pt(9)
d.add_paragraph("_" * 75)

d.add_paragraph("Date: 24/09/2026")
d.add_paragraph(f"To\nThe Purchase Department (Kind Attn: {BUYER['buyer_name']} Madam)\n{BUYER['name']}\nSIPCOT Phase II, Hosur")
d.add_paragraph(f"Respected Madam,\n\nSub: Our best offer for corrugated boxes – your RFQ No. {BUYER['rfq_no']}")

paras = [
    "Greetings from Vijay Box Works. We are a Hosur based corrugated box manufacturer for the last 9 years and we are very happy to receive your enquiry. "
    "Since our factory is only 6 km from your plant, we can give quick delivery and better service than outside suppliers. Please find our rates below. "
    "All rates are per box / per piece and GST 18% is extra.",

    f"For the 3 ply boxes, kettle box (230x190x250) will be {r(1)}, steam iron box {r(2)}, pressure cooker 3 litre box {r(8)} and toaster box {r(15)}. "
    f"Plain spares boxes we can do at {r(16)} for small size and {r(17)} for the medium size, these are our regular sizes so rate is very competitive. "
    f"For the printed hair dryer box (4 colour), rate will be Rs. 7.20 to 7.80 depending on final artwork and number of colours.",

    f"For 5 ply items, mixer 500W box {r(4)}, mixer 750W box {r(5)}, mixer master carton for 2 units {r(6)}, induction cooktop box {r(7)}, "
    f"pressure cooker 5L box {r(9)}, fan blade box {r(12)}, fan motor box {r(13)}, table fan box {r(14)} and the large e-commerce outer {r(18)}. "
    f"For RO purifier box, we suggest to use 18 BF paper instead of 22 BF since it will give same performance in our experience, rate will be Rs. 49.80 only. "
    f"Corner pad 150x150 (5 ply) {r(27)} per piece.",

    f"For 7 ply air cooler box rate is {r(11)} and the base tray for air cooler {r(30)}. "
    "Sorry madam, we are not able to quote for the export master carton (4 units) as we do not have export certification at present, "
    "and also the single face rolls we are not manufacturing.",

    f"For die cut items, mailer box (4 colour printed) {r(20)}, mixer jar insert tray {r(23)} and induction insert {r(24)}. "
    "For both partitions (12 cell and 6 cell) we can give Rs. 6.40 each. "
    f"Layer pad 800x600 will be {r(26)}. For the 1200x1000 big layer pad, rate on request as paper size is non standard for our machine.",

    "Commercial terms: Delivery free of cost to your Hosur plant (within 25 km). Printing plate charges Rs. 2,500 per colour per design and die "
    "charges Rs. 8,000 per die will be charged one time. Payment 45 days please, as we are a small unit (MSME). Delivery within 7 days from PO. "
    "This offer is valid for 7 days.",

    f"We can supply your full annual quantity, for example kettle box {indian(I[1]['annual_qty'])} nos and steam iron box {indian(I[2]['annual_qty'])} nos, "
    "there is no problem, our capacity is 600 MT per month and we are presently running at around 65%.",
]
for t in paras:
    d.add_paragraph(t)

d.add_paragraph().add_run("Regarding your questionnaire:").bold = True
qa = [
    "ISO 9001 – we have applied and certification audit is scheduled in November 2026. We will share certificate after that.",
    "Test reports – we will share the test reports later along with samples.",
    "Paper we use is 100% recycled kraft from Tamil Nadu mills. FSC not available.",
    "Printing – we have 2 colour printer slotter in-house; 4 colour printing we get done from our associate unit in Hosur.",
    "We are supplying to 3 auto component companies and one fan company in Hosur for the last 5 years.",
    "Rejected lots we will replace immediately.",
]
for t in qa:
    d.add_paragraph(t, style="List Bullet")

d.add_paragraph("We request you to kindly give us one opportunity and we assure you of best quality and service.\n\nThanking you,\nYours faithfully,\n\nfor VIJAY BOX WORKS\n\nK. Vijayakumar\nProprietor")
d.save(f"{OUT}/Vijay_Box_Works_Quotation.docx")

email = f"""From: Vijay Box Works <vijayboxworks@gmail.com>
To: {BUYER['buyer_email']}
Date: Thu, 24 Sep 2026 11:05:44 +0530
Subject: quotation
Attachments: Vijay_Box_Works_Quotation.docx

Madam pls find attached our quotation. Kindly check and revert.

Thanks
Vijayakumar
94433 50219
Sent from my Android phone
"""
open(f"{OUT}/email.eml", "w").write(email)
print("C done")
